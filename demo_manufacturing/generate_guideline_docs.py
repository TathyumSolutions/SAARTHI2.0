#!/usr/bin/env python3
"""
generate_guideline_docs.py
===========================
Generates the "unstructured document" side of the manufacturing demo: a
small set of internal factory policy documents (.docx) written in-house as
practitioner-style summaries of well-known, publicly discussed
manufacturing/industrial themes (ISO 9001 quality management, occupational
health & safety, supplier/vendor management, equipment maintenance &
calibration, environmental compliance & waste management). These are
original summaries for demo/testing purposes, NOT a reproduction of any
official ISO/regulatory publication, and should not be treated as a
compliance reference - only as sample content for exercising Saarthi's
document-RAG pipeline against realistic-looking policy text.

Usage:
    python generate_guideline_docs.py
"""
from __future__ import annotations

import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output", "guidelines")

NAVY = RGBColor(0x1F, 0x38, 0x64)


def add_title(doc: Document, title: str, doc_code: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run(f"Internal Policy Reference | {doc_code} | PrecisionForge Manufacturing Ltd. (synthetic demo entity)")
    sub_run.italic = True
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run(
        "This is a synthetic, practitioner-style summary prepared for software demo/testing purposes. "
        "It is not an official ISO/regulatory publication and must not be used as a compliance reference."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1):
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str):
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()


# ---------------------------------------------------------------------
# Document 1: Quality Management Policy
# ---------------------------------------------------------------------

def build_quality_management_policy() -> Document:
    doc = Document()
    add_title(doc, "Quality Management Policy", "QMS-001 v3.0")

    add_heading(doc, "1. Purpose and Applicability")
    add_body(doc,
        "This policy sets out PrecisionForge's quality management approach across incoming material "
        "inspection, in-process production checks, and final inspection before dispatch, in line with "
        "the spirit of ISO 9001's process-approach requirements. It applies to every plant, every "
        "product line, and every supplier feeding raw material into production.")

    add_heading(doc, "2. Quality Inspection Stages")
    add_table(doc, ["Stage", "When Performed", "Owner"], [
        ["Incoming Material", "On receipt of a purchase order shipment, before it enters stores", "Quality Inspector"],
        ["In-Process", "At defined checkpoints during the production run", "Quality Inspector / Production Supervisor"],
        ["Final Inspection", "Before the finished goods are released for dispatch", "Quality Inspector"],
        ["Customer Return", "On receipt of a customer-returned lot", "Quality Inspector / Plant Manager"],
    ])

    add_heading(doc, "3. Sampling and Acceptance Criteria")
    add_bullets(doc, [
        "Sample size for each inspection lot is drawn per the plant's approved AQL (Acceptable "
        "Quality Level) sampling plan, scaled to lot size.",
        "A lot with a defect rate below 2.0% is accepted (Pass); between 2.0% and 6.0% is sent for "
        "100% sorting/rework (Rework); above 6.0% is rejected outright (Fail) and the root cause is "
        "logged against the responsible machine/operator/supplier.",
        "Three consecutive Fail results against the same supplier's material triggers a mandatory "
        "supplier quality audit under the Vendor Management Policy (VMP-003).",
    ])

    add_heading(doc, "4. Non-Conformance and Corrective Action")
    add_body(doc,
        "Every Fail or Rework result is logged as a Non-Conformance Report (NCR). The responsible "
        "Production Supervisor determines containment action (quarantine the affected lot, halt the "
        "line if the cause is machine-related) within 4 working hours, and a Root Cause Corrective "
        "Action (RCCA) using the 5-Why or Fishbone method is closed out within 10 working days.")

    add_heading(doc, "5. Calibration of Measuring Equipment")
    add_body(doc,
        "All gauges, calipers, and testing rigs used for quality inspection are calibrated against a "
        "traceable reference standard per the schedule in the Equipment Maintenance & Calibration "
        "Policy (EMP-004); equipment outside its calibration due date is tagged 'Do Not Use' and "
        "removed from the shop floor until recalibrated.")

    add_heading(doc, "6. Management Review")
    add_body(doc,
        "Plant-level quality metrics (defect rate trend, NCR closure TAT, customer complaint volume) "
        "are reviewed monthly by the Plant Manager and consolidated quarterly for the Board's Quality "
        "Committee review.")

    return doc


# ---------------------------------------------------------------------
# Document 2: EHS (Environment, Health & Safety) Policy
# ---------------------------------------------------------------------

def build_ehs_safety_policy() -> Document:
    doc = Document()
    add_title(doc, "Environment, Health and Safety (EHS) Policy", "EHS-002 v2.5")

    add_heading(doc, "1. Commitment")
    add_body(doc,
        "PrecisionForge is committed to providing a safe working environment for every employee, "
        "contractor, and visitor across its plants, consistent with the spirit of the Factories Act "
        "and applicable occupational health and safety standards (OHSAS 18001 / ISO 45001 practice).")

    add_heading(doc, "2. Personal Protective Equipment (PPE)")
    add_bullets(doc, [
        "Safety helmet, safety shoes, and high-visibility vest are mandatory on the shop floor at all "
        "times for every employee, contractor, and visitor.",
        "Machine operators on CNC, press, and welding stations additionally wear cut-resistant gloves, "
        "safety goggles/face shields, and hearing protection as specified for that machine type.",
        "PPE is issued free of cost and replaced on a defined cycle or immediately on visible damage.",
    ])

    add_heading(doc, "3. Machine Safety and Lockout-Tagout (LOTO)")
    add_body(doc,
        "Before any maintenance, cleaning, or unjamming activity on a machine, the Maintenance "
        "Technician isolates its energy source and applies a lockout tag; the machine is not "
        "re-energised until the tag is removed by the same technician who applied it. Every machine "
        "carries a visible emergency stop and interlocked safety guards on moving parts.")

    add_heading(doc, "4. Incident Reporting and Investigation")
    add_table(doc, ["Severity", "Examples", "Reporting TAT", "Investigation Owner"], [
        ["Near Miss", "Slip without fall, tool drop with no contact", "Same shift", "Shift Supervisor"],
        ["First Aid Case", "Minor cut, bruise treated on-site", "Within 4 hours", "Safety Officer"],
        ["Lost Time Injury (LTI)", "Injury requiring time off work", "Within 1 hour", "Safety Officer + Plant Manager"],
        ["Major/Fatal", "Serious injury, fatality, major equipment damage", "Immediate", "Plant Manager + Head Office EHS"],
    ])

    add_heading(doc, "5. Fire and Emergency Preparedness")
    add_bullets(doc, [
        "Fire extinguishers, hydrants, and smoke detectors are inspected monthly and serviced annually "
        "by a certified vendor.",
        "A full-plant fire/evacuation mock drill is conducted at least twice a year, with attendance "
        "and response-time logged.",
        "Material Safety Data Sheets (MSDS) for all hazardous chemicals (paints, solvents, lubricants) "
        "are maintained at the point of use and at the plant safety office.",
    ])

    add_heading(doc, "6. Training")
    add_body(doc,
        "Every new employee and contractor completes a safety induction before floor access is granted, "
        "and machine operators complete equipment-specific safety training before being certified to "
        "operate that machine class independently.")

    return doc


# ---------------------------------------------------------------------
# Document 3: Vendor / Supplier Management Policy
# ---------------------------------------------------------------------

def build_vendor_management_policy() -> Document:
    doc = Document()
    add_title(doc, "Vendor and Supplier Management Policy", "VMP-003 v2.0")

    add_heading(doc, "1. Purpose")
    add_body(doc,
        "This policy governs how PrecisionForge selects, empanels, evaluates, and exits raw-material "
        "and services suppliers, so that the supply base reliably meets cost, quality, and delivery "
        "requirements across the plant network.")

    add_heading(doc, "2. Supplier Empanelment")
    add_bullets(doc, [
        "A prospective supplier submits company registration, GSTIN, quality certifications (ISO 9001 "
        "or equivalent where applicable), and sample material for first-article inspection before "
        "empanelment.",
        "Empanelment is approved by the Plant Manager for purchase orders below the delegated financial "
        "threshold, and by the Head Office Procurement Committee above it.",
        "Every empanelled supplier is assigned a home plant for primary servicing but may be used by "
        "other plants subject to logistics feasibility.",
    ])

    add_heading(doc, "3. Ongoing Performance Evaluation")
    add_body(doc,
        "Suppliers are scored quarterly on the criteria and weightage defined in the Vendor Rating Card "
        "(on-time delivery, incoming quality acceptance rate, price competitiveness, responsiveness), "
        "producing the quality_rating figure carried on the supplier master record. Ratings below the "
        "'Acceptable' threshold trigger a formal improvement plan; two consecutive quarters below "
        "threshold move the supplier to Suspended status pending re-audit.")

    add_heading(doc, "4. Purchase Order Governance")
    add_bullets(doc, [
        "A purchase order is raised only against a rate agreed with an empanelled, active supplier; "
        "urgent off-contract purchases require Plant Manager approval and are capped at a defined "
        "annual value per supplier.",
        "A purchase order is auto-flagged for review if its unit price exceeds the last three accepted "
        "purchase orders for the same material by more than 15%.",
        "Rejected or cancelled purchase orders record a reason code, reviewed monthly to spot systemic "
        "supplier or specification issues.",
    ])

    add_heading(doc, "5. Payment Terms")
    add_body(doc,
        "Standard payment terms range from 15 to 90 days depending on the supplier's rating tier and "
        "category, per the Vendor Rating Card; suppliers under review or suspended are moved to advance "
        "payment only until re-audited and reinstated.")

    add_heading(doc, "6. Exit and Debarment")
    add_body(doc,
        "A supplier is debarred for confirmed quality fraud (falsified test certificates), repeated "
        "safety non-compliance at their facility on audit, or breach of confidentiality/IP terms, and "
        "is not re-empanelled for a minimum of 24 months.")

    return doc


# ---------------------------------------------------------------------
# Document 4: Equipment Maintenance & Calibration Policy
# ---------------------------------------------------------------------

def build_equipment_maintenance_policy() -> Document:
    doc = Document()
    add_title(doc, "Equipment Maintenance and Calibration Policy", "EMP-004 v1.8")

    add_heading(doc, "1. Purpose")
    add_body(doc,
        "This policy defines how PrecisionForge keeps production machinery and measuring equipment "
        "reliable, safe, and within specification, balancing preventive maintenance cost against "
        "unplanned breakdown risk across the machine fleet.")

    add_heading(doc, "2. Maintenance Categories")
    add_table(doc, ["Category", "Trigger", "Typical Response Time"], [
        ["Preventive Maintenance (PM)", "Scheduled per the machine-type PM calendar", "Planned, scheduled downtime"],
        ["Minor Breakdown", "Machine stoppage resolvable on-shift by plant maintenance staff", "Within 4 hours"],
        ["Major Breakdown", "Requires spare-part replacement or specialist/OEM support", "Within 48 hours"],
        ["Extended Outage", "Major structural failure or awaited long-lead spare part", "Case-by-case recovery plan"],
    ])

    add_heading(doc, "3. Preventive Maintenance Scheduling")
    add_body(doc,
        "Every machine follows the PM frequency, lubrication interval, and overhaul interval defined "
        "for its machine type in the Machine Maintenance Schedule reference; a machine skipping two "
        "consecutive scheduled PM cycles is automatically flagged to the Plant Manager for priority "
        "rescheduling.")

    add_heading(doc, "4. Breakdown Classification and Criticality")
    add_bullets(doc, [
        "Every downtime event is classified by category (Scheduled Maintenance / Minor Breakdown / "
        "Major Breakdown / Extended Outage) and criticality (Low / Medium / High / Critical) based on "
        "downtime hours and production impact.",
        "A machine logging 3 or more High/Critical breakdown events within a rolling 90-day window is "
        "escalated for a full condition assessment and, where justified, replacement rather than "
        "continued repair.",
        "Maintenance cost per event is tracked against the machine's book value to flag machines "
        "approaching end-of-economic-life.",
    ])

    add_heading(doc, "5. Calibration of Measuring and Test Equipment")
    add_body(doc,
        "Gauges, calipers, torque wrenches, and testing rigs used in quality inspection are calibrated "
        "against a traceable reference standard at the interval fixed for that instrument class (never "
        "exceeding 12 months), and carry a visible calibration-due label; equipment overdue for "
        "calibration is withdrawn from use immediately.")

    add_heading(doc, "6. Spare Parts Inventory")
    add_body(doc,
        "Critical spares for high-criticality machines are stocked at each plant per a minimum/maximum "
        "level set by Plant Engineering; a stock-out of a critical spare during a breakdown is logged "
        "as a maintenance-planning non-conformance.")

    return doc


# ---------------------------------------------------------------------
# Document 5: Environmental Compliance & Waste Management Policy
# ---------------------------------------------------------------------

def build_environmental_compliance_policy() -> Document:
    doc = Document()
    add_title(doc, "Environmental Compliance and Waste Management Policy", "ENV-005 v1.6")

    add_heading(doc, "1. Purpose and Scope")
    add_body(doc,
        "This policy summarises how PrecisionForge manages its environmental footprint - effluent, "
        "emissions, and industrial waste - across its plants, aligned in spirit with the Environment "
        "(Protection) Act framework and state Pollution Control Board consent-to-operate conditions.")

    add_heading(doc, "2. Waste Classification and Handling")
    add_table(doc, ["Waste Category", "Examples", "Handling"], [
        ["Hazardous Waste", "Used lubricant/cutting oil, paint sludge, solvent residue", "Stored in labelled containers in a bunded area; disposed via an authorised hazardous-waste handler"],
        ["Scrap Metal", "Turnings, off-cuts, rejected castings", "Segregated by alloy and sold to authorised scrap recyclers"],
        ["Plastic Waste", "Molding sprues, rejected plastic parts", "Reground and reused in-process where specification allows, else sent to an authorised recycler"],
        ["E-Waste", "Damaged PCB modules, control panel electronics", "Collected separately and disposed via an authorised e-waste recycler"],
        ["General/Packaging Waste", "Cartons, wooden pallets, general refuse", "Segregated as recyclable/non-recyclable per municipal norms"],
    ])

    add_heading(doc, "3. Emissions and Effluent Monitoring")
    add_bullets(doc, [
        "Stack emissions from furnaces and powder-coating curing ovens are monitored per the frequency "
        "in the plant's Pollution Control Board consent, with a third-party stack test at least annually.",
        "Effluent from surface-treatment/plating processes is treated at the plant's Effluent Treatment "
        "Plant (ETP) before discharge, and ETP outlet parameters are logged daily.",
        "Any exceedance of a consented emission or discharge limit is reported to the Plant Manager "
        "immediately and to the Head Office EHS function within 24 hours.",
    ])

    add_heading(doc, "4. Energy and Resource Efficiency")
    add_body(doc,
        "Each plant tracks specific energy consumption (units per tonne/unit produced) monthly against "
        "a year-on-year reduction target, and water withdrawal is metered separately for process use, "
        "cooling, and domestic use to identify conservation opportunities.")

    add_heading(doc, "5. Regulatory Consents and Renewals")
    add_body(doc,
        "Consent-to-Operate, Consent-to-Establish, and hazardous-waste authorisation renewals are "
        "tracked on a compliance calendar by the Plant Manager, with a renewal application filed at "
        "least 90 days before expiry.")

    add_heading(doc, "6. Reporting")
    add_body(doc,
        "A consolidated environmental performance report - waste generated/disposed by category, "
        "emissions test results, and any regulatory non-compliance - is placed before the Board's "
        "Sustainability Committee every quarter.")

    return doc


DOCUMENTS = [
    ("quality_management_policy.docx", build_quality_management_policy),
    ("ehs_safety_policy.docx", build_ehs_safety_policy),
    ("vendor_management_policy.docx", build_vendor_management_policy),
    ("equipment_maintenance_policy.docx", build_equipment_maintenance_policy),
    ("environmental_compliance_policy.docx", build_environmental_compliance_policy),
]


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    for filename, builder in DOCUMENTS:
        doc = builder()
        path = os.path.join(OUTPUT_DIR, filename)
        doc.save(path)
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
