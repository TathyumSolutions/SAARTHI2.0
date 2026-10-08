"""
Tests for chart readability rules (top-5 bars, single color, titles,
truncated labels, no misleading pie) and the full-result Excel export
offered when the chat only shows part of a result.
"""
import io
import os
import sys

import pytest
from flask import Flask
from openpyxl import load_workbook

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from app.services.databridge_services.agents.data_visualizer_agent import DataVisualizerAgent
from app.services import result_export_service as exports


def _viz():
    return DataVisualizerAgent()


def _branch_rows(n=110):
    # The grouped "employees by branch" shape after the join fix: one row
    # per branch, so branch_name is unique per row.
    return [{"branch_id": i, "branch_name": f"Branch {i:03d}", "count": (i % 7) + 1} for i in range(1, n + 1)]


# ---------------- chart readability ----------------

def test_many_categories_chart_only_top_5_sorted_descending():
    rows = _branch_rows()
    cfg = _viz().generate_multiple_chart_configs(rows, ["branch_id", "branch_name", "count"])

    assert cfg["recommended"] == "bar"
    bar = cfg["bar"]
    assert len(bar["data"]["labels"]) == 5
    values = bar["data"]["datasets"][0]["data"]
    assert values == sorted(values, reverse=True)
    assert cfg["truncated"] is True
    assert cfg["total_categories"] == 110
    assert "top 5 of 110" in cfg["note"]
    # Grouped by the readable name, not the id, and no "Others" bar.
    assert all(l.startswith("Branch ") for l in bar["data"]["labels"])
    assert "Others" not in bar["data"]["labels"]


def test_grouped_unique_name_is_not_mistaken_for_an_id():
    """Previously a one-row-per-branch result made branch_name look like an
    id (unique per row), so the chart fell back to 'Row 1..Row N'."""
    cfg = _viz().generate_multiple_chart_configs(_branch_rows(), ["branch_id", "branch_name", "count"])
    assert not any(l.startswith("Row ") for l in cfg["bar"]["data"]["labels"])


def test_ties_are_cut_stably_by_label():
    rows = [{"branch_name": f"B{i}", "count": 9} for i in (5, 3, 9, 1, 7, 2, 8)]
    cfg = _viz().generate_multiple_chart_configs(rows, ["branch_name", "count"])
    assert cfg["bar"]["data"]["labels"] == ["B1", "B2", "B3", "B5", "B7"]


def test_bar_chart_is_single_color_with_titles_and_no_legend():
    bar = _viz().generate_multiple_chart_configs(_branch_rows(), ["branch_id", "branch_name", "count"])["bar"]
    dataset = bar["data"]["datasets"][0]
    assert isinstance(dataset["backgroundColor"], str)
    opts = bar["options"]
    assert opts["plugins"]["legend"]["display"] is False
    assert opts["plugins"]["title"]["text"] == "Top 5 Branch Name by Count"
    value_axis = opts["scales"]["x"] if opts["indexAxis"] == "y" else opts["scales"]["y"]
    assert value_axis["title"]["text"] == "Count"
    assert value_axis["beginAtZero"] is True
    assert value_axis["ticks"]["precision"] == 0  # whole-number counts


def test_long_labels_are_shortened_and_go_horizontal():
    rows = [{"customer_name": "An Extremely Long Customer Name Private Limited " + str(i), "amount": i * 10.5}
            for i in range(1, 4)]
    bar = _viz().generate_multiple_chart_configs(rows, ["customer_name", "amount"])["bar"]
    assert bar["options"]["indexAxis"] == "y"
    assert all(len(l) <= DataVisualizerAgent.MAX_LABEL_CHARS for l in bar["data"]["labels"])
    assert bar["full_labels"][0].startswith("An Extremely Long Customer Name Private Limited")


def test_pie_only_for_a_complete_small_set():
    small = [{"region": r, "amount": v} for r, v in [("N", 10), ("S", 20), ("E", 30)]]
    cfg = _viz().generate_multiple_chart_configs(small, ["region", "amount"])
    assert cfg["pie"] and cfg["truncated"] is False

    cfg = _viz().generate_multiple_chart_configs(_branch_rows(), ["branch_id", "branch_name", "count"])
    assert cfg["pie"] == {}


def test_time_series_keeps_every_period():
    rows = [{"order_date": f"2026-{m:02d}-01", "amount": m * 100} for m in range(1, 13)]
    cfg = _viz().generate_multiple_chart_configs(rows, ["order_date", "amount"])
    assert cfg["recommended"] == "line"
    assert len(cfg["line"]["data"]["labels"]) == 12
    assert cfg["truncated"] is False


# ---------------- Excel export ----------------

@pytest.fixture
def app_ctx(tmp_path):
    app = Flask(__name__, instance_path=str(tmp_path))
    with app.app_context():
        yield tmp_path


def test_export_round_trip_to_xlsx(app_ctx):
    rows = _branch_rows()
    info = exports.save_result_export(rows, user_id=7, question="Number of employees by branch?")
    assert info["row_count"] == 110

    export = exports.load_result_export(info["id"], 7)
    wb = load_workbook(io.BytesIO(exports.build_result_workbook(export).getvalue()))
    ws = wb["Results"]
    assert [c.value for c in ws[1]] == ["branch_id", "branch_name", "count"]
    assert ws.max_row == 111  # header + every record, not just the top 5
    assert ws["B2"].value == "Branch 001"
    assert wb["About"]["B1"].value == "Number of employees by branch?"
    assert exports.export_filename(export) == "saarthi_Number_of_employees_by_branch.xlsx"


def test_export_is_private_to_its_user(app_ctx):
    info = exports.save_result_export(_branch_rows(3), user_id=7)
    assert exports.load_result_export(info["id"], 7) is not None
    assert exports.load_result_export(info["id"], 8) is None
    assert exports.load_result_export(info["id"], "7") is not None  # JWT identity is a string


def test_export_rejects_path_like_ids(app_ctx):
    assert exports.load_result_export("../../etc/passwd", 7) is None
    assert exports.load_result_export("x" * 32, 7) is None


def test_single_row_or_non_list_is_not_exported(app_ctx):
    assert exports.save_result_export([{"total": 5}], user_id=1) is None
    assert exports.save_result_export("not rows", user_id=1) is None


def test_formula_like_values_are_written_as_text(app_ctx):
    rows = [{"name": "=HYPERLINK(\"http://evil\")", "v": 1}, {"name": "ok", "v": 2}]
    info = exports.save_result_export(rows, user_id=1)
    ws = load_workbook(io.BytesIO(exports.build_result_workbook(exports.load_result_export(info["id"], 1)).getvalue()))["Results"]
    assert ws["A2"].data_type == "s"


def test_expired_exports_are_pruned(app_ctx):
    old = exports.save_result_export(_branch_rows(3), user_id=1)
    path = os.path.join(str(app_ctx), "result_exports", f"{old['id']}.json")
    stale = os.path.getmtime(path) - (exports.EXPORT_TTL_DAYS + 1) * 86400
    os.utime(path, (stale, stale))

    exports.save_result_export(_branch_rows(3), user_id=1)
    assert not os.path.exists(path)
