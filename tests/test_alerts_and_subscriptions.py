"""
Alerts & Subscriptions: rule suggestions and evaluation, schedules,
recipients and email content, the scheduled runner, the API, and run_type
on the answer pipeline.

Runs against in-memory SQLite with a small Flask app (only the alerts
blueprint), and never calls an LLM, a database connection or SMTP - the
saved-query replay and send_email are patched where a test needs them.
"""
import io
import json
import os
import sys
from datetime import datetime, timedelta
from unittest.mock import MagicMock, patch

import pytest
from flask import Flask
from flask_jwt_extended import JWTManager, create_access_token
from openpyxl import load_workbook

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)
# router_service builds OpenAI clients at import time; no request is made.
os.environ.setdefault("OPENAI_API_KEY", "sk-test")

from app import db, limiter  # noqa: E402
from app.models.query_log import QueryLog  # noqa: E402
from app.models.scheduled_job import ScheduledJob, ScheduledJobRun  # noqa: E402
from app.models.user import User  # noqa: E402
from app.services import alert_rules, job_notifications, job_schedule, scheduled_job_runner  # noqa: E402

SALES_ROWS = [
    {"region": "APAC", "sales_amount": 120, "budget_amount": 100},
    {"region": "EU", "sales_amount": 80, "budget_amount": 100},
    {"region": "LATAM", "sales_amount": "50", "budget_amount": 70},
]


# ------------------------------------------------------------------
# Fixtures
# ------------------------------------------------------------------
@pytest.fixture
def app():
    from app.routes import alert_routes

    flask_app = Flask(__name__, template_folder=os.path.join(ROOT_DIR, "app", "templates"))
    flask_app.config.update(
        TESTING=True,
        SQLALCHEMY_DATABASE_URI="sqlite://",
        SQLALCHEMY_BINDS={"core": "sqlite://", "resources": "sqlite://", "workspace": "sqlite://"},
        JWT_SECRET_KEY="test-secret-key-that-is-long-enough-for-hs256",
        RATELIMIT_ENABLED=False,
        RATELIMIT_STORAGE_URI="memory://",
    )
    db.init_app(flask_app)
    JWTManager(flask_app)
    limiter.init_app(flask_app)
    flask_app.register_blueprint(alert_routes.bp)
    with flask_app.app_context():
        db.create_all()
        yield flask_app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def people(app):
    """Two users and an admin at ACME, one user at another company."""
    def make(name, email, company, role="user"):
        user = User(name=name, email=email, company_code=company, role=role, status="active")
        user.set_password("pw")
        db.session.add(user)
        return user
    users = {
        "asha": make("Asha", "asha@acme.com", "ACME"),
        "ravi": make("Ravi", "ravi@acme.co.in", "ACME"),
        "admin": make("Admin", "admin@acme.com", "ACME", role="admin"),
        "zed": make("Zed", "zed@other.com", "OTHER", role="admin"),
    }
    db.session.commit()
    return users


def _query(user, code="QUERY00001", decision="DB", main_query="SELECT region, sales_amount, budget_amount FROM s"):
    row = QueryLog(query_code=code, user_id=user.id, company_code=user.company_code,
                   question="Sales vs budget by region", answer="3 rows", router_decision=decision,
                   main_query=main_query, sources=["sales<Database>"])
    db.session.add(row)
    db.session.commit()
    return row


def _headers(user):
    return {"Authorization": f"Bearer {create_access_token(identity=str(user.id))}"}


def _alert_payload(**overrides):
    payload = {
        "job_type": "A",
        "name": "Regions below budget",
        "source_query_code": "QUERY00001",
        "rule": {"kind": "compare", "column": "sales_amount", "operator": "lt", "compare_column": "budget_amount"},
        "schedule": {"frequency": "daily", "time": "09:00", "timezone": "Asia/Kolkata"},
        "min_gap_minutes": 0,
        "notify_once_until_recovered": False,
        "recipients": "asha@acme.com, ravi@acme.co.in, partner@gmail.com",
        "subject_template": "Saarthi alert: {region} sales below budget",
        "message_template": "{region} sales were {sales_amount}, {gap_percent} below {budget_amount}.",
        "attach_excel": True,
        "attach_scope": "triggered",
    }
    payload.update(overrides)
    return payload


def _replay_ok(rows):
    return patch.object(scheduled_job_runner, "replay_saved_query",
                        return_value={"rows": rows, "answer": "ok", "error": None, "sources": []})


def _make_job(user, **overrides):
    fields = dict(
        job_code="ALRT00001", job_type="A", user_id=user.id, company_code=user.company_code,
        name="Regions below budget", status="active", source_query_code="QUERY00001",
        question="Sales vs budget by region", router_decision="DB", main_query="SELECT 1",
        rule={"kind": "compare", "column": "sales_amount", "operator": "lt",
              "compare_column": "budget_amount", "scope": "each"},
        schedule={"frequency": "daily", "time": "09:00", "timezone": "UTC"},
        next_run_at=datetime(2026, 10, 21, 9, 0),
        recipients=["asha@acme.com"], subject_template="Alert {name}",
        message_template="{triggered_list}", attach_excel=True, attach_scope="triggered",
    )
    fields.update(overrides)
    job = ScheduledJob(**fields)
    db.session.add(job)
    db.session.commit()
    return job


# ------------------------------------------------------------------
# Rules
# ------------------------------------------------------------------
def test_suggestions_pair_actual_with_budget_and_end_with_any_rows():
    result = alert_rules.suggest_rules(SALES_ROWS)
    kinds = [s["rule"]["kind"] for s in result["suggestions"]]
    assert result["label_column"] == "region"
    assert result["numeric_columns"] == ["sales_amount", "budget_amount"]
    assert kinds == ["compare", "pct_vs_column", "pct_change", "threshold", "row_count"]
    first = result["suggestions"][0]
    assert first["label"] == "sales_amount falls below budget_amount"
    assert first["rule"] == {"kind": "compare", "column": "sales_amount", "operator": "lt",
                             "compare_column": "budget_amount", "scope": "each"}


def test_suggestions_watch_cost_columns_for_going_up():
    rows = [{"plant": "P1", "scrap_cost": 10}, {"plant": "P2", "scrap_cost": 30}]
    first = alert_rules.suggest_rules(rows)["suggestions"][0]
    assert first["rule"]["kind"] == "threshold"
    assert first["rule"]["operator"] == "gt"
    assert "above" in first["label"]


def test_suggestions_skip_id_and_year_columns():
    rows = [{"id": 1, "year": 2026, "orders": 5}, {"id": 2, "year": 2026, "orders": 7}]
    assert alert_rules.suggest_rules(rows)["numeric_columns"] == ["orders"]


def test_compare_rule_fires_for_each_row_below_budget():
    rule = alert_rules.normalize_rule({"kind": "compare", "column": "sales_amount", "operator": "lt",
                                       "compare_column": "budget_amount"})
    result = alert_rules.evaluate_rule(rule, SALES_ROWS)
    assert result["fired"] is True
    assert result["triggered_labels"] == ["EU", "LATAM"]
    assert result["summary"] == "2 of 3 rows met the rule: EU, LATAM."
    assert result["first"]["_gap_percent"] == -20.0


def test_pct_vs_column_only_counts_gaps_bigger_than_the_percent():
    rule = {"kind": "pct_vs_column", "column": "sales_amount", "compare_column": "budget_amount",
            "direction": "below", "percent": 25, "scope": "each"}
    assert alert_rules.evaluate_rule(rule, SALES_ROWS)["triggered_labels"] == ["LATAM"]


def test_pct_change_compares_with_the_last_run_and_not_on_the_first():
    rule = {"kind": "pct_change", "column": "sales_amount", "direction": "drop", "percent": 15, "scope": "each"}
    first = alert_rules.evaluate_rule(rule, SALES_ROWS)
    assert first["fired"] is False
    assert first["summary"].startswith("First run")
    later_rows = [dict(r) for r in SALES_ROWS]
    later_rows[0]["sales_amount"] = 90  # APAC 120 -> 90, a 25% drop
    second = alert_rules.evaluate_rule(rule, later_rows, first["values"])
    assert second["triggered_labels"] == ["APAC"]


def test_threshold_on_the_total_and_row_filter():
    total = {"kind": "threshold", "column": "sales_amount", "operator": "lt", "value": 300, "scope": "total"}
    assert alert_rules.evaluate_rule(total, SALES_ROWS)["fired"] is True  # 250 < 300
    apac_only = {"kind": "threshold", "column": "sales_amount", "operator": "lt", "value": 100,
                 "scope": "each", "row_filter": {"column": "region", "value": "apac"}}
    assert alert_rules.evaluate_rule(apac_only, SALES_ROWS)["fired"] is False


def test_row_count_rule():
    rule = {"kind": "row_count", "operator": "gt", "value": 0}
    assert alert_rules.evaluate_rule(rule, [])["fired"] is False
    assert alert_rules.evaluate_rule(rule, SALES_ROWS)["fired"] is True


@pytest.mark.parametrize("rule, message", [
    (None, "Choose when"),
    ({"kind": "threshold", "column": "x", "operator": "lt"}, "Type the amount"),
    ({"kind": "pct_change", "column": "x", "percent": 0}, "percentage above 0"),
    ({"kind": "compare", "column": "x", "operator": "lt"}, "compare against"),
])
def test_normalize_rule_explains_what_is_missing(rule, message):
    with pytest.raises(ValueError, match=message):
        alert_rules.normalize_rule(rule)


def test_normalize_rule_checks_columns_exist():
    with pytest.raises(ValueError, match="no column named profit"):
        alert_rules.normalize_rule({"kind": "threshold", "column": "profit", "operator": "lt", "value": 1},
                                   ["region", "sales_amount"])


@pytest.mark.parametrize("text, expected", [
    ("tell me if APAC sales goes below 5 lakh",
     {"kind": "threshold", "column": "sales_amount", "operator": "lt", "value": 500000.0,
      "row_filter": {"column": "region", "value": "APAC"}}),
    ("any region more than 10% below budget",
     {"kind": "pct_vs_column", "column": "sales_amount", "compare_column": "budget_amount",
      "direction": "below", "percent": 10.0}),
    ("sales drops more than 15%", {"kind": "pct_change", "column": "sales_amount", "direction": "drop", "percent": 15.0}),
    ("sales above budget", {"kind": "compare", "column": "sales_amount", "operator": "gt", "compare_column": "budget_amount"}),
    ("total sales under 2 crore", {"kind": "threshold", "scope": "total", "value": 20000000.0}),
    ("if any rows come back", {"kind": "row_count", "operator": "gt"}),
])
def test_parse_rule_text_reads_common_phrasings_without_ai(text, expected):
    llm = MagicMock()
    rule = alert_rules.parse_rule_text(text, SALES_ROWS, llm_invoke=llm)
    llm.assert_not_called()
    for key, value in expected.items():
        assert rule[key] == value, key
    assert rule["text"] == text


def test_parse_rule_text_asks_the_llm_only_when_the_parser_cannot():
    llm = MagicMock(return_value='Sure: {"kind": "threshold", "column": "budget_amount", "operator": "gte", "value": 90}')
    rule = alert_rules.parse_rule_text("whenever the plan looks too ambitious", SALES_ROWS, llm_invoke=llm)
    llm.assert_called_once()
    assert rule["kind"] == "threshold" and rule["column"] == "budget_amount" and rule["operator"] == "gte"


def test_parse_rule_text_rejects_llm_output_naming_unknown_columns():
    llm = MagicMock(return_value='{"kind": "threshold", "column": "profit", "operator": "lt", "value": 1}')
    assert alert_rules.parse_rule_text("whenever things look off", SALES_ROWS, llm_invoke=llm) is None


def test_describe_rule():
    rule = {"kind": "pct_vs_column", "column": "sales_amount", "compare_column": "budget_amount",
            "direction": "below", "percent": 10, "scope": "each"}
    assert alert_rules.describe_rule(rule, "region") == "sales_amount is more than 10% below budget_amount, for any region"


# ------------------------------------------------------------------
# Schedules
# ------------------------------------------------------------------
def test_daily_schedule_in_the_jobs_time_zone():
    schedule = job_schedule.normalize_schedule({"frequency": "daily", "time": "9:00", "timezone": "Asia/Kolkata"})
    # 09:00 IST is 03:30 UTC.
    assert job_schedule.next_run_after(schedule, datetime(2026, 10, 21, 2, 0)) == datetime(2026, 10, 21, 3, 30)
    assert job_schedule.next_run_after(schedule, datetime(2026, 10, 21, 3, 30)) == datetime(2026, 10, 22, 3, 30)


def test_weekly_hourly_and_monthly_schedules():
    weekly = job_schedule.normalize_schedule({"frequency": "weekly", "weekday": 0, "time": "08:00", "timezone": "UTC"})
    assert job_schedule.next_run_after(weekly, datetime(2026, 10, 21, 9, 0)) == datetime(2026, 10, 26, 8, 0)
    hourly = job_schedule.normalize_schedule({"frequency": "hourly", "minute": 15, "timezone": "UTC"})
    assert job_schedule.next_run_after(hourly, datetime(2026, 10, 21, 9, 20)) == datetime(2026, 10, 21, 10, 15)
    monthly = job_schedule.normalize_schedule({"frequency": "monthly", "day_of_month": 31, "time": "09:00", "timezone": "UTC"})
    # February has no 31st: run on its last day.
    assert job_schedule.next_run_after(monthly, datetime(2026, 2, 10)) == datetime(2026, 2, 28, 9, 0)
    assert job_schedule.describe_schedule(monthly) == "Monthly, 31st at 09:00 (UTC)"


@pytest.mark.parametrize("raw, message", [
    ({"frequency": "yearly"}, "how often"),
    ({"frequency": "daily", "time": "nine"}, "HH:MM"),
    ({"frequency": "daily", "timezone": "Mars/Base"}, "time zone"),
    ({"frequency": "hourly", "minute": 75}, "minute"),
])
def test_bad_schedules_are_rejected(raw, message):
    with pytest.raises(ValueError, match=message):
        job_schedule.normalize_schedule(raw)


# ------------------------------------------------------------------
# Recipients and the email
# ------------------------------------------------------------------
def test_recipients_can_be_any_address(people):
    asha = people["asha"]
    assert job_notifications.clean_recipients(asha, "me; Ravi@ACME.co.in, someone@gmail.com, me") == [
        "asha@acme.com", "ravi@acme.co.in", "someone@gmail.com"]
    with pytest.raises(ValueError, match="not an email address"):
        job_notifications.clean_recipients(asha, ["not-an-email"])
    with pytest.raises(ValueError, match="at least one"):
        job_notifications.clean_recipients(asha, "")


def test_alert_email_fills_in_values_and_attaches_only_triggered_rows(people):
    job = ScheduledJob(
        job_type="A", name="Regions below budget", question="Sales vs budget <by region>",
        source_query_code="QUERY00001", schedule={"frequency": "daily", "time": "09:00", "timezone": "UTC"},
        rule={"kind": "compare", "column": "sales_amount", "operator": "lt", "compare_column": "budget_amount", "scope": "each"},
        subject_template="Saarthi alert: {region} sales below budget",
        message_template="{region} sales were {sales_amount}, {gap_percent} below {budget_amount}. {unknown}",
        attach_excel=True, attach_scope="triggered", include_table=True,
    )
    evaluation = alert_rules.evaluate_rule(job.rule, SALES_ROWS)
    email = job_notifications.build_email(job, SALES_ROWS, evaluation, datetime(2026, 10, 21, 9, 0))
    assert email["subject"] == "Saarthi alert: EU sales below budget"
    assert "EU sales were 80, 20.0% below 100. {unknown}" in email["text"]
    assert "&lt;by region&gt;" in email["html"]  # escaped
    assert email["attachment_name"] == "Regions_below_budget_2026-10-21.xlsx"
    sheet = load_workbook(io.BytesIO(email["attachment"][1]))["Results"]
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    assert values[0] == ["region", "sales_amount", "budget_amount"]
    assert [v[0] for v in values[1:]] == ["EU", "LATAM"]


def test_subscription_email_uses_full_result_and_marks_tests(people):
    job = ScheduledJob(job_type="S", name="Weekly sales", question="q", schedule={"frequency": "daily", "timezone": "UTC"},
                       subject_template="", message_template="", attach_excel=False, include_table=True)
    email = job_notifications.build_email(job, SALES_ROWS, None, datetime(2026, 10, 21), is_test=True)
    assert email["subject"] == "[Test] Saarthi: Weekly sales for 21 Oct 2026"
    assert "It has 3 rows." in email["text"]
    assert email["attachment"] is None
    assert "This is a test" in email["html"]


# ------------------------------------------------------------------
# The runner
# ------------------------------------------------------------------
def test_alert_run_sends_when_met_and_records_the_run(people):
    job = _make_job(people["asha"])
    with _replay_ok(SALES_ROWS), \
         patch.object(scheduled_job_runner, "_log_run_query", return_value="QUERY00009"), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)) as send:
        run = scheduled_job_runner.execute_job(job, now=datetime(2026, 10, 21, 9, 0))
    assert run.status == "sent"
    assert run.query_code == "QUERY00009"
    assert run.triggered_count == 2
    assert run.notified_to == ["asha@acme.com"]
    to, subject, text, html, attachments = send.call_args[0]
    assert to == ["asha@acme.com"] and subject == "Alert Regions below budget" and text.startswith("EU, LATAM")
    assert attachments[0][0].endswith(".xlsx")
    assert job.is_triggered is True and job.last_fired_at == datetime(2026, 10, 21, 9, 0)
    assert job.last_value == {"APAC": 120.0, "EU": 80.0, "LATAM": 50.0}


def test_once_until_recovered_holds_repeat_alerts(people):
    job = _make_job(people["asha"], notify_once_until_recovered=True)
    healthy = [{"region": "EU", "sales_amount": 120, "budget_amount": 100}]
    statuses = []
    with patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)) as send:
        for rows in (SALES_ROWS, SALES_ROWS, healthy, SALES_ROWS):
            with _replay_ok(rows):
                statuses.append(scheduled_job_runner.execute_job(job).status)
    assert statuses == ["sent", "fired", "ok", "sent"]
    assert send.call_count == 2


def test_minimum_gap_holds_alerts_inside_it(people):
    job = _make_job(people["asha"], min_gap_minutes=60 * 24)
    start = datetime(2026, 10, 21, 9, 0)
    with _replay_ok(SALES_ROWS), patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)):
        first = scheduled_job_runner.execute_job(job, now=start)
        second = scheduled_job_runner.execute_job(job, now=start + timedelta(hours=1))
        third = scheduled_job_runner.execute_job(job, now=start + timedelta(days=1, minutes=1))
    assert [first.status, second.status, third.status] == ["sent", "fired", "sent"]
    assert "minimum gap" in second.value


def test_failed_run_marks_error_and_emails_only_the_owner_once(people):
    job = _make_job(people["asha"], recipients=["ravi@acme.co.in"])
    failing = patch.object(scheduled_job_runner, "replay_saved_query",
                           return_value={"rows": [], "answer": "", "error": "Database unreachable"})
    with failing, patch.object(scheduled_job_runner, "send_email", return_value=(True, None)) as send:
        scheduled_job_runner.execute_job(job)
        run = scheduled_job_runner.execute_job(job)
    assert run.status == "failed" and run.error == "Database unreachable"
    assert job.status == "error"
    assert send.call_count == 1
    assert send.call_args[0][0] == ["asha@acme.com"]
    with _replay_ok(SALES_ROWS), patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)):
        scheduled_job_runner.execute_job(job)
    assert job.status == "active" and job.last_error is None


def test_subscription_skips_empty_results_when_asked(people):
    job = _make_job(people["asha"], job_code="SUBS00001", job_type="S", rule=None, skip_if_empty=True)
    with _replay_ok([]), patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email") as send:
        run = scheduled_job_runner.execute_job(job)
    assert run.status == "skipped"
    send.assert_not_called()


def test_run_due_jobs_claims_each_due_job_once_and_skips_paused(people):
    asha = people["asha"]
    due = _make_job(asha)
    paused = _make_job(asha, job_code="ALRT00002", status="paused")
    later = _make_job(asha, job_code="ALRT00003", next_run_at=datetime(2026, 10, 22, 9, 0))
    now = datetime(2026, 10, 21, 9, 0, 30)
    with _replay_ok(SALES_ROWS), patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)):
        runs = scheduled_job_runner.run_due_jobs(now=now)
        again = scheduled_job_runner.run_due_jobs(now=now)
    assert [r.job_id for r in runs] == [due.id]
    assert again == []
    assert db.session.get(ScheduledJob, due.id).next_run_at == datetime(2026, 10, 22, 9, 0)
    assert ScheduledJobRun.query.filter_by(job_id=paused.id).count() == 0
    assert ScheduledJobRun.query.filter_by(job_id=later.id).count() == 0


def test_inactive_owner_fails_the_run_without_querying(people):
    asha = people["asha"]
    job = _make_job(asha)
    asha.status = "pending"
    db.session.commit()
    with patch.object(scheduled_job_runner, "replay_saved_query") as replay, \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)):
        run = scheduled_job_runner.execute_job(job)
    replay.assert_not_called()
    assert run.status == "failed" and "not active" in run.error


def test_replay_uses_the_liked_query_replay_with_no_insights(people):
    from app.services import router_service

    job = _make_job(people["asha"])
    fake = MagicMock(return_value={"error": False, "table": [{"a": 1}], "answer": "1 row", "sources": []})
    with patch.dict(router_service.REUSE_DISPATCH, {"DB": fake}):
        result = scheduled_job_runner.replay_saved_query(job)
    assert result == {"rows": [{"a": 1}], "answer": "1 row", "error": None, "sources": []}
    question, matched, score, ctx = fake.call_args[0]
    assert matched.main_query == "SELECT 1" and matched.query_code == "QUERY00001"
    assert ctx["user_id"] == people["asha"].id
    assert ctx["with_insights"] is False


def test_replay_refuses_tracks_that_need_ai(people):
    job = _make_job(people["asha"], router_decision="FILES")
    assert "database and spreadsheet" in scheduled_job_runner.replay_saved_query(job)["error"]


def test_scheduled_run_logs_query_with_run_type_and_job(people):
    job = _make_job(people["asha"])
    replay = {"rows": SALES_ROWS, "answer": "3 rows", "error": None, "sources": ["sales<Database>"]}
    code = scheduled_job_runner._log_run_query(job, replay)
    row = QueryLog.query.filter_by(query_code=code).one()
    assert row.run_type == "A"
    assert row.scheduled_job_id == job.id
    assert row.execution_type == "reused" and row.matched_query_code == "QUERY00001"


def test_reused_db_query_passes_with_insights_through(people):
    from app.services import router_service
    from app.services.databridge_services import langgraph_agent

    matched = MagicMock(query_code="Q1", main_query="SELECT 1", sources=[], question="q", remarks=None)
    with patch.object(langgraph_agent.query_formatter, "execute_query",
                      return_value={"status": "success", "data": [], "message": "ok"}) as execute, \
         patch.object(langgraph_agent, "_build_db_config_for_user", return_value=None):
        router_service._run_reused_db_query("q", matched, 1.0, {"session_id": "s", "with_insights": False})
    assert execute.call_args.kwargs["with_insights"] is False


def test_get_smart_response_stores_run_type(people):
    from app.services import router_service

    with patch.object(router_service, "classify_query_heuristic", return_value="GENERAL"), \
         patch.object(router_service, "answer_general_knowledge", return_value={"answer": "hi", "steps": []}), \
         patch.object(router_service, "_log_query", return_value="QUERY00002") as log:
        router_service.RouterService().get_smart_response("hello", user_id=1, run_type="S")
        router_service.RouterService().get_smart_response("hello", user_id=1, run_type="bogus")
    assert log.call_args_list[0].kwargs["run_type"] == "S"
    assert log.call_args_list[1].kwargs["run_type"] == "Q"


def _call(name, question, tables):
    return {"name": name, "args": {"question": question, "tables": tables}, "id": question}


def test_scheduling_merges_per_table_database_calls_into_one():
    from app.services import router_service

    question = "Inspections and production orders with defect rate above 5%"
    calls = [_call("query_database", "inspections", ["inspection"]),
             _call("query_database", "orders", ["production_order", "inspection"])]
    merged = router_service._merge_calls_for_scheduling(calls, question)
    assert len(merged) == 1
    assert merged[0]["name"] == "query_database"
    assert merged[0]["args"] == {"question": question, "tables": ["inspection", "production_order"]}


def test_scheduling_leaves_mixed_or_single_calls_alone():
    from app.services import router_service

    mixed = [_call("query_database", "a", ["t1"]), _call("query_spreadsheet_data", "b", ["s1"])]
    assert router_service._merge_calls_for_scheduling(mixed, "q") == mixed
    single = [_call("query_database", "a", ["t1"])]
    assert router_service._merge_calls_for_scheduling(single, "q") == single
    docs = [{"name": "search_documents", "args": {"question": "a", "document_codes": []}},
            {"name": "search_documents", "args": {"question": "b", "document_codes": ["D1"]}}]
    assert router_service._merge_calls_for_scheduling(docs, "q") == docs


def test_router_prompt_asks_for_one_query_only_when_scheduling():
    from app.services import router_service

    def prompt(scheduled):
        with patch.object(router_service, "_log_build_router_messages_params"), \
             patch.object(router_service, "_log_strategy_event"):
            return router_service._build_router_messages(
                "q", [], {}, "", "", scheduled_setup=scheduled)[0].content
    assert "SCHEDULED QUESTION" in prompt(True)
    assert "SCHEDULED QUESTION" not in prompt(False)


@pytest.mark.parametrize("decision, words", [
    ("MULTI", "more than one data source"),
    ("FILES", "files and documents"),
    ("GENERAL", "without looking up your data"),
    ("SOMETHING_NEW", "another kind of source"),
])
def test_unsupported_answers_say_where_they_came_from(app, people, decision, words):
    _query(people["asha"], decision=decision)
    res = app.test_client().post("/api/alerts", json=_alert_payload(), headers=_headers(people["asha"]))
    assert res.status_code == 400 and words in res.json["error"]


# ------------------------------------------------------------------
# The API
# ------------------------------------------------------------------
def test_create_alert_from_own_query(app, people):
    asha = people["asha"]
    _query(asha)
    client = app.test_client()
    res = client.post("/api/alerts", json=_alert_payload(), headers=_headers(asha))
    assert res.status_code == 201, res.json
    job = res.json["job"]
    assert job["job_code"] == "ALRT00001"
    assert job["main_query"].startswith("SELECT region")
    assert job["recipients"] == ["asha@acme.com", "ravi@acme.co.in", "partner@gmail.com"]
    assert job["rule_text"] == "sales_amount is below budget_amount"
    assert job["next_run_at"].endswith("Z")
    sub = client.post("/api/alerts", json=_alert_payload(job_type="S", rule=None), headers=_headers(asha))
    assert sub.json["job"]["job_code"] == "SUBS00001"
    assert sub.json["job"]["rule"] is None


@pytest.mark.parametrize("change, message", [
    ({"recipients": "not-an-email"}, "not an email address"),
    ({"rule": None}, "Choose when"),
    ({"name": ""}, "Give it a name"),
    ({"source_query_code": "QUERY00404"}, "not found"),
])
def test_create_rejects_bad_forms(app, people, change, message):
    asha = people["asha"]
    _query(asha)
    res = app.test_client().post("/api/alerts", json=_alert_payload(**change), headers=_headers(asha))
    assert res.status_code == 400
    assert message in res.json["error"]


def test_create_rejects_someone_elses_query_and_non_replayable_answers(app, people):
    _query(people["ravi"], code="QUERY00001")
    _query(people["asha"], code="QUERY00002", decision="FILES")
    client = app.test_client()
    res = client.post("/api/alerts", json=_alert_payload(), headers=_headers(people["asha"]))
    assert res.status_code == 400 and "not found" in res.json["error"]
    res = client.post("/api/alerts", json=_alert_payload(source_query_code="QUERY00002"), headers=_headers(people["asha"]))
    assert res.status_code == 400 and "database or spreadsheet" in res.json["error"]


def test_visibility_owner_admin_and_other_company(app, people):
    asha, ravi, admin, zed = people["asha"], people["ravi"], people["admin"], people["zed"]
    job = _make_job(asha)
    _make_job(ravi, job_code="ALRT00002")
    client = app.test_client()
    assert [j["job_code"] for j in client.get("/api/alerts/", headers=_headers(asha)).json["jobs"]] == ["ALRT00001"]
    admin_list = client.get("/api/alerts/", headers=_headers(admin)).json
    assert admin_list["is_company_admin"] is True
    assert sorted(j["job_code"] for j in admin_list["jobs"]) == ["ALRT00001", "ALRT00002"]
    assert client.get(f"/api/alerts/{job.id}", headers=_headers(ravi)).status_code == 404
    assert client.get(f"/api/alerts/{job.id}", headers=_headers(zed)).status_code == 404
    assert client.post(f"/api/alerts/{job.id}/pause", headers=_headers(admin)).json["job"]["status"] == "paused"


def test_pause_resume_duplicate_delete(app, people):
    asha = people["asha"]
    job = _make_job(asha, status="error", last_error="boom")
    client = app.test_client()
    resumed = client.post(f"/api/alerts/{job.id}/resume", headers=_headers(asha)).json["job"]
    assert resumed["status"] == "active" and resumed["last_error"] is None
    copy = client.post(f"/api/alerts/{job.id}/duplicate", headers=_headers(asha))
    assert copy.status_code == 201
    assert copy.json["job"]["status"] == "paused" and copy.json["job"]["name"] == "Copy of Regions below budget"
    assert client.delete(f"/api/alerts/{job.id}", headers=_headers(asha)).status_code == 200
    assert db.session.get(ScheduledJob, job.id) is None


def test_update_keeps_the_copied_query_and_owner(app, people):
    asha, admin = people["asha"], people["admin"]
    _query(asha)
    client = app.test_client()
    job_id = client.post("/api/alerts", json=_alert_payload(), headers=_headers(asha)).json["job"]["id"]
    QueryLog.query.filter_by(query_code="QUERY00001").delete()
    db.session.commit()
    res = client.put(f"/api/alerts/{job_id}", json=_alert_payload(name="Renamed"), headers=_headers(admin))
    assert res.status_code == 200, res.json
    assert res.json["job"]["name"] == "Renamed"
    assert res.json["job"]["user_id"] == asha.id
    assert res.json["job"]["main_query"].startswith("SELECT region")


def test_queries_picker_lists_only_own_replayable_queries(app, people):
    asha = people["asha"]
    _query(asha, code="QUERY00001")
    _query(asha, code="QUERY00002", decision="API")
    _query(people["ravi"], code="QUERY00003")
    codes = [q["query_code"] for q in app.test_client().get("/api/alerts/queries", headers=_headers(asha)).json["queries"]]
    assert codes == ["QUERY00001"]


def test_source_runs_the_saved_query_and_suggests_rules(app, people):
    asha = people["asha"]
    _query(asha)
    from app.routes import alert_routes
    with patch.object(alert_routes, "replay_saved_query",
                      return_value={"rows": SALES_ROWS, "answer": "", "error": None}):
        res = app.test_client().post("/api/alerts/source", json={"query_code": "QUERY00001", "job_type": "A"},
                                     headers=_headers(asha))
    assert res.status_code == 200
    assert res.json["supported"] is True and res.json["row_count"] == 3
    assert res.json["suggestions"][0]["label"] == "sales_amount falls below budget_amount"


def test_rule_endpoints(app, people):
    client = app.test_client()
    headers = _headers(people["asha"])
    parsed = client.post("/api/alerts/rules/parse", json={"text": "sales below budget", "rows": SALES_ROWS}, headers=headers)
    assert parsed.status_code == 200
    assert parsed.json["sentence"] == "sales_amount is below budget_amount, for any region"
    assert parsed.json["evaluation"]["triggered_labels"] == ["EU", "LATAM"]
    evaluated = client.post("/api/alerts/rules/evaluate", json={
        "rule": {"kind": "threshold", "column": "sales_amount", "operator": "gt", "value": 100}, "rows": SALES_ROWS,
    }, headers=headers)
    assert evaluated.json["evaluation"]["summary"] == "1 of 3 rows met the rule: APAC."


def test_preview_and_test_send(app, people):
    asha = people["asha"]
    _query(asha)
    client = app.test_client()
    with _replay_ok(SALES_ROWS):
        preview = client.post("/api/alerts/preview", json=_alert_payload(), headers=_headers(asha))
    assert preview.status_code == 200, preview.json
    assert preview.json["subject"] == "Saarthi alert: EU sales below budget"
    assert preview.json["would_send"] is True and preview.json["triggered_count"] == 2
    assert ScheduledJob.query.count() == 0  # nothing saved
    with _replay_ok(SALES_ROWS), patch.object(scheduled_job_runner, "send_email", return_value=(True, None)) as send:
        sent = client.post("/api/alerts/test", json=_alert_payload(), headers=_headers(asha))
    assert sent.status_code == 200
    assert send.call_args[0][0] == ["asha@acme.com"]  # the creator only, not the recipients
    assert send.call_args[0][1].startswith("[Test]")


def test_run_now_records_a_manual_run(app, people):
    asha = people["asha"]
    job = _make_job(asha)
    with _replay_ok(SALES_ROWS), patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)):
        res = app.test_client().post(f"/api/alerts/{job.id}/run", headers=_headers(asha))
    assert res.json["run"]["trigger"] == "manual" and res.json["run"]["status"] == "sent"
    runs = app.test_client().get(f"/api/alerts/{job.id}/runs", headers=_headers(asha)).json["runs"]
    assert len(runs) == 1


def test_settings(app, people):
    data = app.test_client().get("/api/alerts/settings", headers=_headers(people["asha"])).json
    assert data["own_email"] == "asha@acme.com"
    assert data["default_templates"]["A"]["subject"] == "Saarthi alert: {name}"
    assert json.dumps(data)  # serialisable


def test_pages_render(app):
    from app.routes import page_routes
    app.register_blueprint(page_routes.bp)
    client = app.test_client()
    create = client.get("/alerts/new").get_data(as_text=True)
    manage = client.get("/alerts").get_data(as_text=True)
    assert "New alert or subscription" in create and "View sample report" in create and "Send me a test now" in create
    assert "Recent runs" in manage and "/api/alerts/" in manage


def test_run_due_jobs_command(app, people):
    from app.commands import run_due_jobs_command
    job_id = _make_job(people["asha"], next_run_at=datetime.utcnow() - timedelta(minutes=1)).id
    with _replay_ok(SALES_ROWS), patch.object(scheduled_job_runner, "_log_run_query", return_value=None), \
         patch.object(scheduled_job_runner, "send_email", return_value=(True, None)):
        result = app.test_cli_runner().invoke(run_due_jobs_command)
    assert result.exit_code == 0, result.output
    assert f"job {job_id}: sent" in result.output
